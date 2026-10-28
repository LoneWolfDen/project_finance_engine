"""Tests for the Excel project planner.

Builds the workbook with scripts/build_planner.py, recalculates it with LibreOffice
(headless) and checks the scheduling results against hand-worked dates.

Run:  python -m unittest discover -s excel-project-planner/tests -v
Needs openpyxl and LibreOffice (`soffice` on PATH); recalculation tests are skipped without it.
"""
import datetime as dt
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "scripts" / "build_planner.py"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
STATUS_DATE = dt.date(2026, 10, 25)
ERRORS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")

# Tasks sheet columns (data starts on row 6, so task ID n sits on row n + 5)
COL = {"WBS": "B", "PS": "N", "PF": "O", "FS": "R", "FF": "S", "Var": "T", "Prog": "V", "Status": "W",
       "Work": "Z", "Cost": "AA", "Sprint": "AB", "Pred": "AC"}


def row(task_id):
    return task_id + 5


def build(path):
    subprocess.run([sys.executable, str(BUILD), str(path)], check=True, capture_output=True)


def recalc(src, outdir, edits=None):
    """Optionally edit input cells, then let LibreOffice recalculate and save a copy."""
    wb = openpyxl.load_workbook(src)
    wb["Settings"]["C8"] = STATUS_DATE          # StatusDateOverride -> deterministic results
    for (sheet, ref), value in (edits or {}).items():
        wb[sheet][ref] = value
    staged = Path(outdir) / "in" / src.name
    staged.parent.mkdir(parents=True, exist_ok=True)
    wb.save(staged)
    out = Path(outdir) / "out"
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(staged)],
                   check=True, capture_output=True, timeout=600)
    return openpyxl.load_workbook(out / src.name, data_only=True)


def d(value):
    return value.date() if isinstance(value, dt.datetime) else value


class StructureTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "Project_Planner.xlsx"
        build(cls.path)
        cls.wb = openpyxl.load_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames, ["Instructions", "Dashboard", "Settings", "Tasks", "Gantt", "Resources",
                                              "Sprints", "Holidays", "MSP_Tasks", "MSP_Resources", "Lists"])

    def test_named_ranges(self):
        names = set(self.wb.defined_names)
        for n in ("ProjectStart", "StatusDate", "WorkWeek", "Holidays", "HoursPerDay", "GanttStart", "GanttPeriod",
                  "Sprint1Start", "SprintLength", "tsk_ID", "tsk_PlanStart", "tsk_FcFinish", "tsk_Resource",
                  "res_Name", "res_Rate", "spr_Name", "spr_Start", "hdr_PlanFinish"):
            self.assertIn(n, names)

    def test_scheduling_formulas_use_names(self):
        f = self.wb["Tasks"]["N7"].value
        for token in ("ProjectStart", "WorkWeek", "Holidays"):
            self.assertIn(token, f)

    def test_msp_headers_match_project_fields(self):
        headers = [c.value for c in self.wb["MSP_Tasks"][1]]
        for h in ("ID", "Name", "Outline Level", "Duration", "Predecessors", "Resource Names", "% Complete",
                  "Actual Start", "Actual Finish", "Baseline Start", "Baseline Finish"):
            self.assertIn(h, headers)

    def test_no_post_2010_functions(self):
        banned = ("XLOOKUP", "FILTER(", "SEQUENCE(", "UNIQUE(", "LET(", "_xlfn.")
        for ws in self.wb.worksheets:
            for r in ws.iter_rows(max_row=12):
                for c in r:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for b in banned:
                            self.assertNotIn(b, c.value, f"{ws.title}!{c.coordinate}")


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ScheduleTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "Project_Planner.xlsx"
        build(cls.path)
        cls.wb = recalc(cls.path, cls.tmp.name)
        cls.ts = cls.wb["Tasks"]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def val(self, task_id, key):
        return d(self.ts[f"{COL[key]}{row(task_id)}"].value)

    def test_no_formula_errors(self):
        bad = [f"{ws.title}!{c.coordinate}={c.value}" for ws in self.wb.worksheets for r in ws.iter_rows() for c in r
               if isinstance(c.value, str) and c.value.startswith(ERRORS)]
        self.assertEqual(bad, [])

    def test_wbs_and_outline(self):
        self.assertEqual(self.val(1, "WBS"), "1")
        self.assertEqual(self.val(13, "WBS"), "3.1.2")
        self.assertEqual(self.val(30, "WBS"), "5.3")

    def test_finish_to_start(self):
        self.assertEqual(self.val(2, "PS"), dt.date(2026, 9, 7))
        self.assertEqual(self.val(2, "PF"), dt.date(2026, 9, 9))
        self.assertEqual(self.val(3, "PS"), dt.date(2026, 9, 10))

    def test_start_to_start_with_lag(self):
        self.assertEqual(self.val(7, "PS"), dt.date(2026, 9, 17))   # task 6 starts Mon 14-Sep, SS+3

    def test_finish_to_finish_with_lag(self):
        self.assertEqual(self.val(18, "PF"), dt.date(2026, 10, 29))  # task 17 finishes 28-Oct, FF+1
        self.assertEqual(self.val(18, "PS"), dt.date(2026, 10, 27))

    def test_two_predecessors_take_latest(self):
        self.assertEqual(self.val(9, "PS"), dt.date(2026, 10, 1))   # max(task 7, task 8 finish)

    def test_weekend_skipped(self):
        self.assertEqual(self.val(6, "PF"), dt.date(2026, 9, 23))   # 8 days from Mon 14-Sep

    def test_holidays_skipped(self):
        # Hypercare: 10 days from Wed 16-Dec skipping 24, 25, 31 Dec and 1 Jan
        self.assertEqual(self.val(26, "PF"), dt.date(2027, 1, 4))

    def test_start_no_earlier_than(self):
        self.assertEqual(self.val(22, "PS"), dt.date(2026, 12, 7))

    def test_manual_task(self):
        self.assertEqual(self.val(28, "PS"), dt.date(2026, 11, 16))

    def test_summary_rollup(self):
        self.assertEqual(self.val(10, "PS"), dt.date(2026, 10, 2))
        self.assertEqual(self.val(10, "PF"), dt.date(2026, 10, 30))
        self.assertEqual(self.val(11, "Work"), 168)   # 40 + 80 + 48 hrs
        self.assertAlmostEqual(self.val(1, "Prog"), 1.0)

    def test_actuals_drive_forecast(self):
        self.assertEqual(self.val(2, "FF"), dt.date(2026, 9, 10))   # actual finish used
        self.assertEqual(self.val(13, "FF"), dt.date(2026, 10, 30))  # in progress, 50% of 10 days left after status
        self.assertEqual(self.val(14, "FS"), dt.date(2026, 11, 2))   # successor pushed
        self.assertEqual(self.val(13, "Status"), "Behind")
        self.assertEqual(self.val(14, "Status"), "Late Start")
        self.assertEqual(self.val(10, "Var"), 6)

    def test_predecessor_text(self):
        self.assertEqual(self.val(7, "Pred"), "6SS+3d")
        self.assertEqual(self.val(17, "Pred"), "16,13SS+2d")
        self.assertEqual(self.val(18, "Pred"), "17FF+1d")

    def test_sprint_and_cost(self):
        self.assertEqual(self.val(2, "Sprint"), "Sprint 1")
        self.assertEqual(self.val(2, "Cost"), 24 * 95)
        self.assertEqual(self.wb["Dashboard"]["C15"].value, "Sprint 4")

    def test_msp_export(self):
        mx = self.wb["MSP_Tasks"]
        self.assertEqual(mx["E7"].value, "8d")                      # task 6 duration
        self.assertEqual(mx["I19"].value, "Casey Kim[50%]")         # task 18 resource units
        self.assertEqual(mx["O23"].value, "Start No Earlier Than")  # task 22 constraint


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class RippleTest(unittest.TestCase):
    """A late actual finish must push every successor's forecast dates."""

    def test_slip_ripples_to_successors(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Project_Planner.xlsx"
            build(path)
            # Task 13 (row 18) actually finishes Wed 04-Nov
            wb = recalc(path, tmp, {("Tasks", "Q18"): dt.date(2026, 11, 4)})
            ts = wb["Tasks"]
            self.assertEqual(d(ts["R19"].value), dt.date(2026, 11, 5))    # task 14 starts next working day
            self.assertEqual(d(ts["S19"].value), dt.date(2026, 11, 12))
            self.assertEqual(d(ts["R24"].value), dt.date(2026, 11, 12))   # task 19 milestone
            self.assertEqual(d(ts["R26"].value), dt.date(2026, 11, 13))   # task 21 testing
            # planned dates are untouched by actuals
            self.assertEqual(d(ts["O19"].value), dt.date(2026, 10, 30))

    def test_status_date_moves_unstarted_work(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Project_Planner.xlsx"
            build(path)
            # remove task 13's actual start: it is now unstarted and must begin on/after the status date
            wb = recalc(path, tmp, {("Tasks", "P18"): None, ("Tasks", "U18"): None})
            ts = wb["Tasks"]
            self.assertEqual(d(ts["R18"].value), dt.date(2026, 10, 26))   # first working day after 25-Oct
            self.assertEqual(d(ts["S18"].value), dt.date(2026, 11, 6))
            self.assertEqual(d(ts["R19"].value), dt.date(2026, 11, 9))    # task 14 follows


if __name__ == "__main__":
    unittest.main()
