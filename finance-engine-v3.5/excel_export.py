"""
Builds the Excel workbook for POST /api/export/xlsx with openpyxl.

The browser does all the calculations and sends a spec; this module only lays it out:

{
  "currency_symbol": "£",                 # used by "money" columns without their own symbol
  "sheets": [{
    "name": "Burndown",
    "title": "PO Burndown (GBP)",          # optional line above the table
    "notes": ["..."],                      # optional lines under the title
    "header": ["Month", "Forecast Spend"], # table header (row after title/notes)
    "rows": [["Jul-25", 1234.5], ...],    # a cell may be {"value": 5, "symbol": "₹"} for its own currency
    "formats": ["text", "money"],          # per column: text|money|money2|int|pct|date
    "symbols": [null, "₹"],                # optional per-column currency symbol
    "widths": [10, 16],                    # optional column widths
    "hidden": false,
    "charts": [{
      "type": "bar" | "line",
      "title": "...", "y_title": "...", "y2_title": "...",
      "series": [{"col": 2, "type": "bar"|"line", "secondary": false, "color": "2196F3", "dash": false}],
      "anchor": "H2", "width": 24, "height": 12
    }]
  }]
}
Columns in "series" are 1-based; the first column holds the categories.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(italic=True, color="666666", size=9)


def number_format(kind, symbol):
    sym = '"' + symbol.replace('"', '') + '"' if symbol else ""
    return {
        "money": f"{sym}#,##0;-{sym}#,##0",
        "money2": f"{sym}#,##0.00;-{sym}#,##0.00",
        "int": "#,##0",
        "pct": "0%",
    }.get(kind)


def add_chart(ws, spec, header_row, first_row, last_row):
    """Native chart over the table: categories in column 1, series from spec['series']."""
    if last_row < first_row:
        return
    cats = Reference(ws, min_col=1, min_row=first_row, max_row=last_row)
    primary = BarChart() if spec.get("type", "bar") == "bar" else LineChart()
    secondary = None
    for s in spec["series"]:
        kind = s.get("type", spec.get("type", "bar"))
        target = primary
        if s.get("secondary") or (kind == "line") != isinstance(primary, LineChart):
            if secondary is None:
                secondary = LineChart() if kind == "line" else BarChart()
            target = secondary
        target.add_data(Reference(ws, min_col=s["col"], min_row=header_row, max_row=last_row), titles_from_data=True)
        series = target.series[-1]
        color = s.get("color")
        if kind == "line":
            series.smooth = False
            if color:
                series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 22000
            if s.get("dash"):
                series.graphicalProperties.line.dashStyle = "dash"
        elif color:
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
    primary.set_categories(cats)
    primary.title = spec.get("title")
    primary.y_axis.title = spec.get("y_title")
    primary.legend.position = "b"
    primary.width = spec.get("width", 24)
    primary.height = spec.get("height", 12)
    # openpyxl hides axes unless delete is set to False explicitly
    primary.x_axis.delete = False
    primary.y_axis.delete = False
    primary.x_axis.tickLblPos = "low"  # month labels below the plot, also when values are negative
    if secondary is not None:
        secondary.set_categories(cats)
        if spec.get("y2_title") or any(s.get("secondary") for s in spec["series"]):
            secondary.y_axis.axId = 200
            secondary.y_axis.title = spec.get("y2_title")
            secondary.y_axis.crosses = "max"
            secondary.y_axis.delete = False
            secondary.y_axis.majorGridlines = None
        primary += secondary
    ws.add_chart(primary, spec.get("anchor", "H2"))


def write_sheet(wb, spec, default_symbol):
    ws = wb.create_sheet(spec["name"][:31])
    row = 1
    if spec.get("title"):
        ws.cell(row=row, column=1, value=spec["title"]).font = TITLE_FONT
        row += 1
    for note in spec.get("notes", []):
        ws.cell(row=row, column=1, value=note).font = NOTE_FONT
        row += 1
    if spec.get("title") or spec.get("notes"):
        row += 1
    header = spec.get("header") or []
    header_row = row
    for c, label in enumerate(header, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    first_row = row + 1 if header else row
    formats = spec.get("formats") or []
    symbols = spec.get("symbols") or []
    for r, values in enumerate(spec.get("rows", []), start=first_row):
        for c, value in enumerate(values, start=1):
            own_symbol = None
            if isinstance(value, dict):
                own_symbol, value = value.get("symbol"), value.get("value")
            cell = ws.cell(row=r, column=c, value=value)
            kind = formats[c - 1] if c - 1 < len(formats) else "text"
            symbol = own_symbol or (symbols[c - 1] if c - 1 < len(symbols) else None) or default_symbol
            fmt = number_format(kind, symbol)
            if fmt and isinstance(value, (int, float)):
                cell.number_format = fmt
    last_row = first_row + len(spec.get("rows", [])) - 1
    if header:
        ws.freeze_panes = ws.cell(row=first_row, column=1)
        if last_row >= first_row:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(header))}{last_row}"
    widths = spec.get("widths") or []
    for c in range(1, max(len(header), len(widths)) + 1):
        width = widths[c - 1] if c - 1 < len(widths) and widths[c - 1] else None
        if width is None:
            longest = max([len(str(header[c - 1])) if c - 1 < len(header) else 0] +
                          [len(str(v[c - 1].get("value") if isinstance(v[c - 1], dict) else v[c - 1]))
                           for v in spec.get("rows", [])[:200] if c - 1 < len(v) and v[c - 1] is not None])
            width = min(max(10, longest + 2), 50)
        ws.column_dimensions[get_column_letter(c)].width = width
    for chart in spec.get("charts", []):
        add_chart(ws, chart, header_row, first_row, last_row)
    if spec.get("hidden"):
        ws.sheet_state = "hidden"


def build_workbook(spec):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec.get("sheets", []):
        write_sheet(wb, sheet, spec.get("currency_symbol", ""))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
